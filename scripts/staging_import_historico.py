"""Staging de migracion historica de campanas OCU26 - Etapa 2B.

Lee OCUPACION_2026.xlsx (fuente historica, grilla mensual por elemento) y
OCU26_BASE_DATOS.xlsx (MAESTRO_ELEMENTOS + CAMPANAS) y genera un staging
clasificado (IMPORTAR / YA_EXISTE / REVISAR / NO_IMPORTAR) listo para una
carga posterior a CAMPANAS. NO importa nada.

Estrictamente READ-ONLY sobre ambos Excel: verifica SHA-256 al inicio y al
final. Solo escribe en el directorio --salida.

Uso:
    python scripts/staging_import_historico.py --fuente <OCUPACION_2026.xlsx>
        --salida <dir> [--base <OCU26_BASE_DATOS.xlsx>] [--sha-base <sha256>]

La base se resuelve igual que el resto del pipeline (validate_input.
resolve_input_path). Ninguna ruta personal se versiona.

Reglas:
- Hojas: CENCO F, CENCO D, PLED, REM-PIL, TRIPSTORE Y LS D (solo filas
  TRIPSTORE), AEROPUERTOS F/D. YPF, LS F, CENCOMEDIA no se procesan; filas
  London/LS o Cencomedia dentro de hojas procesadas -> NO_IMPORTAR.
- OT = IDCampana. Prefijos de la celda OT (PUBLI, PORCO, ...) se conservan
  como trazabilidad; "B <n>" (bonificada) nunca se asume IDCampana.
- ElementoOrigen se resuelve contra MAESTRO_ELEMENTOS por evidencia: codigo
  exacto, normalizacion de espacios, prefijo UN->UNI confirmado (nunca replace
  global), sufijo de slot -Vn, y para PLED la Descripcion del maestro.
  Solo confianza ALTA llega a IMPORTAR.
- Fechas: se consolidan rangos INICIO/FIN identicos o contiguos por
  IDCampana + ElementoID; los huecos generan periodos separados; toda
  contradiccion, mes sin presencia o fecha no interpretable -> REVISAR.
- Deduplicacion contra CAMPANAS por IDCampana + ElementoID + fechas.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from validate_input import CAMPANAS_HEADERS, resolve_input_path  # noqa: E402

EXPECTED_BASE_SHA256 = "4e1ff067d25a5cb95b5cb4ce09389f06ece5a65db6537d87f5b414fac6adde2d"
ANIO_ALCANCE = 2026

# Hojas a procesar -> medio esperado (None = mixto, se toma del maestro).
HOJAS_PROCESAR: dict[str, str | None] = {
    "CENCO F": "Estático",
    "CENCO D": "Digital",
    "PLED": "Digital",
    "REM-PIL": None,
    "TRIPSTORE Y LS D": "Digital",
    "AEROPUERTOS F": "Estático",
    "AEROPUERTOS D": "Digital",
}
# Hojas excluidas completamente (no se procesan) -> motivo.
HOJAS_EXCLUIDAS: dict[str, str] = {
    "YPF": "YPF",
    "LS F": "LONDON/LS",
    "CENCOMEDIA": "CENCOMEDIA",
}
# Valores de "Clas" que, dentro de una hoja procesada, quedan fuera de alcance.
CLAS_EXCLUIDAS: dict[str, str] = {
    "LONDON SUPPLY": "LONDON/LS",
    "LONDON": "LONDON/LS",
    "LS": "LONDON/LS",
    "CENCOMEDIA": "CENCOMEDIA",
    "YPF": "YPF",
}
# Sin OT numerica: celda OT o PAUTA con estos textos = uso no comercial
# (decoracion navidena, institucional del soporte/shopping).
PATRON_NO_COMERCIAL = re.compile(r"\bDECO\b|NAVIDE|\bINSTI|INSTITUCIONAL|\bCENCO\b|CENCOSUD|LANDMARK")

ESTADOS = ("IMPORTAR", "YA_EXISTE", "REVISAR", "NO_IMPORTAR")

MESES_ES = {
    "ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6,
    "JUL": 7, "AGO": 8, "SEP": 9, "SET": 9, "OCT": 10, "NOV": 11, "DIC": 12,
}

_TEXTOS_INDET = {"INDET", "INDEF", "INDEFINIDO", "INDEFINIDA", "INDETERMINADO", "INDETERMINADA"}
_EXCEL_EPOCH = dt.date(1899, 12, 30)


class StagingError(Exception):
    pass


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def texto(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        if v.is_integer():
            v = int(v)
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def normalizar_codigo(v: Any) -> str:
    """Mayusculas, espacios colapsados y sin espacios alrededor de '-'."""
    s = texto(v).upper()
    return re.sub(r"\s*-\s*", "-", s)


def compactar_codigo(v: Any) -> str:
    return re.sub(r"\s+", "", texto(v).upper())


def col_letra(idx0: int) -> str:
    return openpyxl.utils.get_column_letter(idx0 + 1)


def inicio_mes(d: dt.date) -> dt.date:
    return d.replace(day=1)


def fin_mes(d: dt.date) -> dt.date:
    nxt = (d.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
    return nxt - dt.timedelta(days=1)


def sumar_meses(d: dt.date, n: int) -> dt.date:
    y, m = divmod(d.year * 12 + (d.month - 1) + n, 12)
    return dt.date(y, m + 1, 1)


def meses_entre(ini: dt.date, fin: dt.date) -> list[dt.date]:
    out, m = [], inicio_mes(ini)
    while m <= fin:
        out.append(m)
        m = sumar_meses(m, 1)
    return out


def fmt_fecha(d: dt.date | None) -> str:
    return d.isoformat() if d else ""


def fmt_meses(meses: Iterable[dt.date]) -> str:
    return ", ".join(m.strftime("%Y-%m") for m in sorted(meses))


# ---------------------------------------------------------------------------
# Parsing de celdas
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class OTToken:
    numero: int
    prefijo: str = ""
    bonificada: bool = False


@dataclass
class CeldaOT:
    raw: str
    tokens: list[OTToken]
    texto_no_ot: str  # texto residual sin numero (DECO NAV, S/OT, SALA VIP, ...)


def parse_ot_cell(value: Any) -> CeldaOT:
    """Separa una celda OT en tokens numericos.

    '2186/4284' -> 2 tokens; 'PUBLI 4554 / B 4726' -> PUBLI 4554 + 4726
    bonificada; 'DECO NAV' -> sin tokens, texto_no_ot='DECO NAV'.
    """
    raw = texto(value)
    if not raw:
        return CeldaOT(raw="", tokens=[], texto_no_ot="")
    tokens: list[OTToken] = []
    residual: list[str] = []
    for parte in re.split(r"/", raw):
        parte = parte.strip()
        if not parte:
            continue
        m = re.fullmatch(r"(.*?)(\d{3,6})\s*", parte)
        if not m:
            residual.append(parte)
            continue
        pref = m.group(1).strip().upper()
        bonif = False
        pm = re.fullmatch(r"(.*?)\s*\bB\b", pref)
        if pm:
            bonif = True
            pref = pm.group(1).strip()
        if pref == "" and tokens:
            # '4381/4400' tras 'PUBLI 4381' conserva el prefijo del token anterior
            pref = tokens[-1].prefijo
        tok = OTToken(numero=int(m.group(2)), prefijo=pref, bonificada=bonif)
        if tok not in tokens:
            tokens.append(tok)
    return CeldaOT(raw=raw, tokens=tokens, texto_no_ot=" / ".join(residual).upper())


def parse_fecha(value: Any) -> tuple[dt.date | None, str]:
    """-> (fecha, tipo) con tipo en FECHA, SERIAL, INDET, VACIA, INVALIDA."""
    if value is None:
        return None, "VACIA"
    if isinstance(value, dt.datetime):
        return value.date(), "FECHA"
    if isinstance(value, dt.date):
        return value, "FECHA"
    if isinstance(value, bool):
        return None, "INVALIDA"
    if isinstance(value, (int, float)):
        if 30000 <= value <= 60000 and float(value).is_integer():
            return _EXCEL_EPOCH + dt.timedelta(days=int(value)), "SERIAL"
        return None, "INVALIDA"
    s = texto(value).upper()
    if not s:
        return None, "VACIA"
    if s in _TEXTOS_INDET:
        return None, "INDET"
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        try:
            return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))), "FECHA"
        except ValueError:
            return None, "INVALIDA"
    return None, "INVALIDA"


# ---------------------------------------------------------------------------
# Estructura de hoja
# ---------------------------------------------------------------------------

_ATTR_HEADERS = {
    "CIUDAD": "Ciudad", "SHOP": "Shop", "VIA COM": "Shop", "CLAS": "Clas", "UBIC": "Ubic",
    "DISP": "Disp", "CODIGO": "Codigo", "NIVEL": "Nivel", "DESCRIPCION": "Descripcion",
    "DETALLE": "Detalle", "MATERIAL": "Material",
}


@dataclass
class Bloque:
    col: int          # indice 0 de la columna OT
    mes: dt.date      # primer dia del mes
    etiqueta: str


@dataclass
class EstructuraHoja:
    fila_header: int  # indice 0
    attr_cols: dict[str, int]
    bloques: list[Bloque]


def detectar_estructura(rows: list[tuple], hoja: str) -> EstructuraHoja:
    for i, r in enumerate(rows[:12]):
        vals = [sin_acentos(texto(v)).upper() for v in r]
        if "PAUTA" in vals and "CODIGO" in vals:
            break
    else:
        raise StagingError(f"{hoja}: no se encontro fila de encabezado (Codigo + PAUTA)")
    header = [sin_acentos(texto(v)).upper() for v in rows[i]]
    meses_row = rows[i - 1] if i > 0 else ()
    cols_ot = [j for j in range(len(header) - 1) if header[j + 1] == "PAUTA"]
    if not cols_ot:
        raise StagingError(f"{hoja}: sin bloques mensuales")
    attr_cols: dict[str, int] = {}
    for j in range(cols_ot[0]):
        nombre = _ATTR_HEADERS.get(header[j])
        if nombre and nombre not in attr_cols:
            attr_cols[nombre] = j
    if "Codigo" not in attr_cols:
        raise StagingError(f"{hoja}: sin columna Codigo")

    etiquetas = [header[j] for j in cols_ot]
    explicitos: list[dt.date | None] = []
    for j in cols_ot:
        v = meses_row[j] if j < len(meses_row) else None
        explicitos.append(inicio_mes(v.date() if isinstance(v, dt.datetime) else v) if isinstance(v, (dt.datetime, dt.date)) else None)
    meses = inferir_meses(etiquetas, explicitos, hoja)
    bloques = [Bloque(col=j, mes=m, etiqueta=e) for j, m, e in zip(cols_ot, meses, etiquetas)]
    return EstructuraHoja(fila_header=i, attr_cols=attr_cols, bloques=bloques)


def inferir_meses(etiquetas: list[str], explicitos: list[dt.date | None], hoja: str) -> list[dt.date]:
    """Completa meses faltantes por secuencia y valida contra la etiqueta."""
    if not any(explicitos):
        raise StagingError(f"{hoja}: fila de meses sin fechas")
    ancla = next(k for k, m in enumerate(explicitos) if m)
    meses = [sumar_meses(explicitos[ancla], k - ancla) for k in range(len(etiquetas))]
    for k, (et, exp, m) in enumerate(zip(etiquetas, explicitos, meses)):
        if exp and exp != m:
            raise StagingError(f"{hoja}: bloque {k} mes {exp} rompe la secuencia mensual ({m})")
        num = MESES_ES.get(et[:3])
        if num is None or num != m.month:
            raise StagingError(f"{hoja}: etiqueta '{et}' no coincide con el mes inferido {m:%Y-%m}")
    return meses


# ---------------------------------------------------------------------------
# Extraccion de ocurrencias
# ---------------------------------------------------------------------------

@dataclass
class Ocurrencia:
    hoja: str
    fila: int               # fila Excel (1-based)
    ref: str                # 'HOJA!K3'
    mes: dt.date
    codigo: str
    attrs: dict[str, str]
    ot_raw: str
    token: OTToken | None
    n_tokens: int
    texto_no_ot: str
    pauta: str
    ini: dt.date | None
    fin: dt.date | None
    tipo_ini: str
    tipo_fin: str
    obs: str
    alcance: str            # EN_ALCANCE o motivo de exclusion

    @property
    def compartida(self) -> bool:
        return self.n_tokens > 1


def alcance_fila(hoja: str, attrs: dict[str, str]) -> str:
    clas = normalizar_codigo(attrs.get("Clas", ""))
    if clas in CLAS_EXCLUIDAS:
        return CLAS_EXCLUIDAS[clas]
    if hoja == "TRIPSTORE Y LS D" and clas != "TRIPSTORE":
        return "LONDON/LS"
    return "EN_ALCANCE"


def extraer_hoja(rows: list[tuple], hoja: str) -> list[Ocurrencia]:
    est = detectar_estructura(rows, hoja)
    out: list[Ocurrencia] = []
    for i in range(est.fila_header + 1, len(rows)):
        r = rows[i]
        get = lambda j: r[j] if j < len(r) else None  # noqa: E731
        attrs = {k: texto(get(j)) for k, j in est.attr_cols.items()}
        codigo = attrs.get("Codigo", "")
        alcance = alcance_fila(hoja, attrs) if codigo else "FILA_SIN_CODIGO"
        for b in est.bloques:
            vals = [get(b.col + k) for k in range(5)]
            if all(texto(v) == "" for v in vals):
                continue
            celda = parse_ot_cell(vals[0])
            ini, tipo_ini = parse_fecha(vals[2])
            fin, tipo_fin = parse_fecha(vals[3])
            base = dict(
                hoja=hoja, fila=i + 1, ref=f"{hoja}!{col_letra(b.col)}{i + 1}", mes=b.mes,
                codigo=codigo, attrs=attrs, ot_raw=celda.raw, n_tokens=len(celda.tokens),
                texto_no_ot=celda.texto_no_ot, pauta=texto(vals[1]), ini=ini, fin=fin,
                tipo_ini=tipo_ini, tipo_fin=tipo_fin, obs=texto(vals[4]), alcance=alcance,
            )
            if celda.tokens:
                for t in celda.tokens:
                    out.append(Ocurrencia(token=t, **base))
            else:
                out.append(Ocurrencia(token=None, **base))
    return out


def leer_fuente(path: Path) -> tuple[list[Ocurrencia], dict[str, list[dt.date]], dict[str, str]]:
    """-> (ocurrencias, ventana de meses por hoja, hojas no procesadas -> motivo)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ocurrencias: list[Ocurrencia] = []
        ventanas: dict[str, list[dt.date]] = {}
        no_procesadas: dict[str, str] = {}
        nombres = {ws.title.strip(): ws.title for ws in wb.worksheets}
        for hoja in HOJAS_PROCESAR:
            if hoja not in nombres:
                raise StagingError(f"Hoja requerida ausente en la fuente: {hoja}")
        for limpio, real in nombres.items():
            if limpio in HOJAS_PROCESAR:
                continue
            no_procesadas[real] = (
                f"EXCLUIDA ({HOJAS_EXCLUIDAS[limpio]})" if limpio in HOJAS_EXCLUIDAS else "NO LISTADA (fuera de alcance)"
            )
        for hoja in HOJAS_PROCESAR:
            ws = wb[nombres[hoja]]
            head = list(ws.iter_rows(min_row=1, max_row=12, values_only=True))
            est = detectar_estructura(head, hoja)
            max_col = est.bloques[-1].col + 5
            rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=max_col, values_only=True))
            ocurrencias.extend(extraer_hoja(rows, hoja))
            ventanas[hoja] = [b.mes for b in est.bloques]
        return ocurrencias, ventanas, no_procesadas
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Crosswalk de elementos
# ---------------------------------------------------------------------------

@dataclass
class Match:
    elemento_id: str = ""
    metodo: str = "SIN_MATCH"
    confianza: str = "BAJA"
    observacion: str = ""
    candidatos: list[str] = field(default_factory=list)


# Circuitos del maestro que nunca pueden ser destino de esta migracion.
CIRCUITOS_FUERA_DE_ALCANCE = {"YPF Digital", "YPF Estático", "London Supply"}


class IndiceMaestro:
    def __init__(self, maestro: pd.DataFrame):
        df = maestro.copy()
        df["ElementoID"] = df["ElementoID"].map(texto)
        df = df[~df["CircuitoDashboard"].map(texto).isin(CIRCUITOS_FUERA_DE_ALCANCE)]
        self.df = df
        self.por_id = {r["ElementoID"]: r for r in df.to_dict("records")}
        self.por_compacto: dict[str, list[str]] = defaultdict(list)
        self.pled_por_desc: dict[str, list[str]] = defaultdict(list)
        for eid, r in self.por_id.items():
            self.por_compacto[compactar_codigo(eid)].append(eid)
            if texto(r.get("CircuitoDashboard")) == "Pantalla Led":
                self.pled_por_desc[compactar_codigo(r.get("Descripcion"))].append(eid)

    def buscar_codigo(self, codigo: str) -> list[str]:
        return sorted(set(self.por_compacto.get(compactar_codigo(codigo), [])))


def _sin_sufijo_slot(codigo: str) -> str | None:
    m = re.fullmatch(r"(.+?)-V\d+", normalizar_codigo(codigo))
    return m.group(1) if m else None


def _tokens_desc(s: str) -> set[str]:
    return {t for t in re.findall(r"[A-Z0-9]+", sin_acentos(texto(s)).upper()) if len(t) > 2}


def similitud_desc(a: str, b: str) -> float:
    ta, tb = _tokens_desc(a), _tokens_desc(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)


def descripcion_canonica(s: Any) -> str:
    return " ".join(re.findall(r"[A-Z0-9]+", sin_acentos(texto(s)).upper()))


def posicion_codigo(codigo: str) -> str:
    """Segmento final de posicion: 'UN-DOBLES-3F-4-V2' -> '3F-4'."""
    base = _sin_sufijo_slot(codigo) or normalizar_codigo(codigo)
    m = re.search(r"(\d+[A-Z]+-\d+)$", base)
    return m.group(1) if m else ""


def candidatos_por_codigo(codigo: str, hoja: str, idx: IndiceMaestro) -> tuple[list[str], str]:
    """Busqueda por evidencia de codigo, en orden de fuerza. -> (candidatos, metodo)."""
    exactos = [e for e in idx.por_id if e == texto(codigo)]
    if exactos:
        return exactos, "EXACTO"
    c = idx.buscar_codigo(codigo)
    if c:
        return c, "NORMALIZACION_ESPACIOS"
    norm = normalizar_codigo(codigo)
    variantes: list[tuple[str, str]] = []
    if norm.startswith("UN-") and not norm.startswith("UNI-"):
        variantes.append(("UNI-" + norm[3:], "PREFIJO_UN_UNI"))
    base = _sin_sufijo_slot(norm)
    if base:
        variantes.append((base, "SUFIJO_SLOT"))
        if base.startswith("UN-") and not base.startswith("UNI-"):
            variantes.append(("UNI-" + base[3:], "SUFIJO_SLOT+PREFIJO_UN_UNI"))
    for v, metodo in variantes:
        c = idx.buscar_codigo(v)
        if c:
            return c, metodo
    if hoja == "PLED":
        raiz = base or norm
        c = sorted(set(idx.pled_por_desc.get(compactar_codigo(raiz), [])))
        if c:
            return c, "PLED_DESCRIPCION_MAESTRO"
    return [], "SIN_MATCH"


def candidatos_por_descripcion(
    hoja: str, codigo: str, attrs: dict[str, str], idx: IndiceMaestro, ubic_alias: dict[str, str]
) -> tuple[list[str], str]:
    """Sin match de codigo: misma Ubicacion + Medio y misma Descripcion.

    - Descripcion identica y misma posicion final del codigo (p.ej. '1G-5'):
      unica -> DESCRIPCION_Y_POSICION (comprobable).
    - Descripcion del maestro contenida al inicio de la de origen (familia):
      varios -> DESCRIPCION_FAMILIA (ambiguo, nunca ALTA).
    """
    ubic = ubic_alias.get(normalizar_codigo(attrs.get("Ubic", "")), normalizar_codigo(attrs.get("Ubic", "")))
    medio = HOJAS_PROCESAR.get(hoja)
    desc = descripcion_canonica(attrs.get("Descripcion", "") or attrs.get("Disp", ""))
    if not desc:
        return [], ""
    pos = posicion_codigo(codigo)
    iguales, familia = [], []
    for eid, r in idx.por_id.items():
        if medio and texto(r.get("Medio")) != medio:
            continue
        if normalizar_codigo(r.get("Ubicacion")) != ubic:
            continue
        dm = descripcion_canonica(r.get("Descripcion"))
        if not dm:
            continue
        if dm == desc:
            iguales.append(eid)
        elif desc.startswith(dm + " "):
            familia.append(eid)
    if pos:
        con_pos = [e for e in iguales if posicion_codigo(e) == pos]
        if len(con_pos) == 1:
            return con_pos, "DESCRIPCION_Y_POSICION"
    if iguales:
        return sorted(iguales), "DESCRIPCION_IGUAL"
    if familia:
        return sorted(familia), "DESCRIPCION_FAMILIA"
    return [], ""


def construir_crosswalk(ocurrencias: list[Ocurrencia], idx: IndiceMaestro) -> pd.DataFrame:
    """Una fila por (HojaOrigen, ElementoOrigen) en alcance con ocurrencias.

    Los elementos de filas fuera de alcance (London/LS, Cencomedia) no se
    mapean: nunca pueden llegar a IMPORTAR.
    """
    grupos: dict[tuple[str, str], list[Ocurrencia]] = defaultdict(list)
    for o in ocurrencias:
        if o.codigo and o.alcance == "EN_ALCANCE":
            grupos[(o.hoja, o.codigo)].append(o)

    # 1) candidatos por codigo
    pre: dict[tuple[str, str], tuple[list[str], str, dict[str, str]]] = {}
    for key, occ in grupos.items():
        attrs = occ[0].attrs
        cands, metodo = candidatos_por_codigo(key[1], key[0], idx)
        pre[key] = (cands, metodo, attrs)

    # 2) mapa aprendido Ubic origen -> Ubicacion maestro (solo matches de codigo directo)
    aprendido: dict[str, Counter] = defaultdict(Counter)
    for (hoja, _), (cands, metodo, attrs) in pre.items():
        if len(cands) == 1 and metodo in ("EXACTO", "NORMALIZACION_ESPACIOS"):
            u = normalizar_codigo(attrs.get("Ubic", ""))
            if u:
                aprendido[u][normalizar_codigo(idx.por_id[cands[0]].get("Ubicacion"))] += 1
    ubic_alias = {u: c.most_common(1)[0][0] for u, c in aprendido.items()}

    # 3) sin match de codigo -> evidencia de ubicacion + descripcion
    for key, (cands, metodo, attrs) in list(pre.items()):
        if not cands:
            c2, m2 = candidatos_por_descripcion(key[0], key[1], attrs, idx, ubic_alias)
            if c2:
                pre[key] = (c2, m2, attrs)

    resueltos: dict[tuple[str, str], tuple[Match, dict[str, str], str, str | None]] = {}
    for (hoja, codigo), (cands, metodo, attrs) in sorted(pre.items()):
        medio_esp = HOJAS_PROCESAR.get(hoja)
        ubic = normalizar_codigo(attrs.get("Ubic", ""))
        desc = attrs.get("Descripcion", "") or attrs.get("Disp", "")
        m = Match(candidatos=cands)
        obs: list[str] = []
        if codigo.upper().startswith("UN-") and "UN_UNI" not in metodo:
            obs.append(f"Prefijo UN- probado como UNI-{normalizar_codigo(codigo)[3:]}: no existe en maestro")
        if len(cands) == 1 and metodo != "DESCRIPCION_FAMILIA":
            eid = cands[0]
            r = idx.por_id[eid]
            m.elemento_id, m.metodo = eid, metodo
            problemas = []
            if medio_esp and texto(r.get("Medio")) != medio_esp:
                problemas.append(f"Medio maestro {texto(r.get('Medio'))} != medio de hoja {medio_esp}")
            ubic_m = normalizar_codigo(r.get("Ubicacion"))
            if ubic and ubic in ubic_alias and ubic_alias[ubic] != ubic_m and aprendido[ubic][ubic_alias[ubic]] >= 2:
                problemas.append(f"Ubicacion maestro {ubic_m} != ubicacion habitual de '{ubic}' ({ubic_alias[ubic]})")
            sim = max(similitud_desc(desc, r.get("Descripcion")), similitud_desc(attrs.get("Disp", ""), r.get("Descripcion")))
            if metodo == "PLED_DESCRIPCION_MAESTRO":
                obs.append(f"Codigo PLED historico '{codigo}' = Descripcion maestro '{texto(r.get('Descripcion'))}'"
                           f" (Ubic origen '{attrs.get('Ubic', '')}' -> {texto(r.get('Ubicacion'))})")
            elif metodo == "DESCRIPCION_Y_POSICION":
                obs.append(f"Codigo historico distinto; misma Ubicacion, Descripcion identica y posicion '{posicion_codigo(codigo)}'"
                           f" ('{codigo}' -> {eid})")
            elif metodo == "DESCRIPCION_IGUAL":
                problemas.append("Solo coincide la Descripcion (sin evidencia de codigo/posicion)")
            elif metodo not in ("EXACTO", "NORMALIZACION_ESPACIOS") and desc and texto(r.get("Descripcion")) and sim == 0:
                problemas.append(f"Descripciones sin palabras en comun ('{desc}' vs '{texto(r.get('Descripcion'))}')")
            if "UN_UNI" in metodo:
                obs.append(f"UN->UNI confirmado contra maestro: '{codigo}' -> {eid}")
            if "SUFIJO_SLOT" in metodo:
                obs.append("Sufijo de slot -Vn removido")
            if metodo == "NORMALIZACION_ESPACIOS":
                obs.append("Diferencia solo de espacios")
            if problemas:
                # Nunca se publica un ElementoID no validado: queda como candidato.
                m.confianza, m.elemento_id = "MEDIA", ""
                obs = problemas + [f"Candidato no confirmado: {eid}"] + obs
            else:
                m.confianza = "ALTA"
        elif cands:
            m.metodo, m.confianza = metodo + ("" if metodo == "DESCRIPCION_FAMILIA" else "_AMBIGUO"), "MEDIA"
            obs.append("Multiples candidatos razonables, sin equivalencia comprobable: " + ", ".join(cands))
        else:
            sugeridos = sugerir_candidatos(hoja, attrs, idx, ubic_alias)
            m.candidatos = [s for s, _ in sugeridos]
            obs.append("Elemento historico inexistente en maestro")
            if sugeridos:
                obs.append("elementos parecidos (no equivalentes): " + ", ".join(f"{s} ({p:.2f})" for s, p in sugeridos))
        m.observacion = " | ".join(obs)
        resueltos[(hoja, codigo)] = (m, attrs, desc, medio_esp)

    heredar_de_slots_hermanos(resueltos)

    filas = []
    for (hoja, codigo), (m, attrs, desc, medio_esp) in sorted(resueltos.items()):
        r = idx.por_id.get(m.elemento_id, {})
        n_filas = len({o.fila for o in grupos[(hoja, codigo)]})
        filas.append({
            "HojaOrigen": hoja,
            "ElementoOrigen": codigo,
            "MedioOrigen": medio_esp or "",
            "CircuitoOrigen": circuito_origen(hoja, attrs),
            "ElementoID_OCU26": m.elemento_id,
            "CircuitoDashboard_OCU26": texto(r.get("CircuitoDashboard")),
            "Subcircuito_OCU26": texto(r.get("Subcircuito")),
            "Ubicacion_OCU26": texto(r.get("Ubicacion")),
            "Descripcion_OCU26": texto(r.get("Descripcion")),
            "MetodoMatch": m.metodo,
            "Confianza": m.confianza,
            "Observacion": m.observacion,
            "UbicacionOrigen": attrs.get("Ubic", ""),
            "DescripcionOrigen": desc,
            "Medio_OCU26": texto(r.get("Medio")),
            "Candidatos": ", ".join(m.candidatos),
            "FilasOrigen": n_filas,
        })
    return pd.DataFrame(filas)


METODOS_POR_CODIGO = ("EXACTO", "NORMALIZACION_ESPACIOS", "PREFIJO_UN_UNI", "SUFIJO_SLOT", "SUFIJO_SLOT+PREFIJO_UN_UNI")


def heredar_de_slots_hermanos(resueltos: dict) -> None:
    """Slots -Vn del mismo codigo base son el mismo elemento fisico.

    Si todos los slots hermanos con confianza ALTA apuntan a un unico
    ElementoID, los slots no resueltos del mismo codigo base lo heredan
    (p.ej. descripciones 'Nivel 4..13' generadas por autocompletado en la
    planilla). No se hereda si la evidencia de codigo propia del slot apunta
    a otro elemento.
    """
    por_base: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    for (hoja, codigo) in resueltos:
        base = _sin_sufijo_slot(codigo)
        if base:
            por_base[(hoja, base)].append((hoja, codigo))
    for (hoja, base), claves in por_base.items():
        altas = {resueltos[k][0].elemento_id for k in claves if resueltos[k][0].confianza == "ALTA"}
        if len(altas) != 1:
            continue
        eid = next(iter(altas))
        for k in claves:
            m = resueltos[k][0]
            if m.confianza == "ALTA":
                continue
            if m.metodo in METODOS_POR_CODIGO and m.candidatos and m.candidatos != [eid]:
                continue
            previo = f"{m.metodo}/{m.confianza}"
            m.elemento_id, m.confianza, m.metodo = eid, "ALTA", "HERENCIA_SLOTS_HERMANOS"
            m.observacion = (f"Slot de '{base}': sus slots hermanos ALTA resuelven a {eid}; "
                             f"evidencia propia del slot descartada ({previo}): {m.observacion}")
            m.candidatos = [eid]


def circuito_origen(hoja: str, attrs: dict[str, str]) -> str:
    clas = texto(attrs.get("Clas", "")).upper()
    if hoja.startswith("AEROPUERTOS"):
        return "AEROPUERTOS"
    if hoja == "PLED":
        return "PLED"
    return clas


def sugerir_candidatos(
    hoja: str, attrs: dict[str, str], idx: IndiceMaestro, ubic_alias: dict[str, str], umbral: float = 0.5
) -> list[tuple[str, float]]:
    """Solo informativo (nunca match): misma Ubicacion y descripcion parecida."""
    u0 = normalizar_codigo(attrs.get("Ubic", ""))
    ubic = ubic_alias.get(u0, u0)
    desc = attrs.get("Descripcion", "") or attrs.get("Disp", "")
    medio = HOJAS_PROCESAR.get(hoja)
    if not desc:
        return []
    out = []
    for eid, r in idx.por_id.items():
        if medio and texto(r.get("Medio")) != medio:
            continue
        if ubic and normalizar_codigo(r.get("Ubicacion")) != ubic:
            continue
        s = similitud_desc(desc, texto(r.get("Descripcion")))
        if s >= umbral:
            out.append((eid, round(s, 2)))
    out.sort(key=lambda x: (-x[1], x[0]))
    return out[:5]


# ---------------------------------------------------------------------------
# Reconstruccion de fechas
# ---------------------------------------------------------------------------

@dataclass
class Periodo:
    ini: dt.date
    fin: dt.date
    regla: str
    pares: list[tuple[dt.date, dt.date]]
    dudas: list[str] = field(default_factory=list)


@dataclass
class ResultadoFechas:
    periodos: list[Periodo]
    presencia: list[dt.date]
    problemas: list[str]  # problemas globales del grupo

    @property
    def disjuntos(self) -> bool:
        return len(self.periodos) > 1


def _intersecta_mes(ini: dt.date, fin: dt.date, mes: dt.date) -> bool:
    return ini <= fin_mes(mes) and fin >= mes


def reconstruir_periodos(ocurrencias: list[Ocurrencia], ventana: Iterable[dt.date]) -> ResultadoFechas:
    """Consolida INICIO/FIN de un grupo IDCampana + elemento.

    - Pares identicos o contiguos (fin + 1 dia = inicio) se consolidan.
    - Pares sin contacto -> periodos separados (huecos reales).
    - Pares solapados distintos -> se unifica el envolvente pero queda en duda.
    - Cada mes con presencia debe caer en un periodo; cada mes de un periodo
      dentro de la ventana de la hoja debe tener presencia.
    - Fechas de celdas con varias OT no se usan (son de la celda, no de la OT).
    """
    ventana = set(ventana)
    presencia = sorted({o.mes for o in ocurrencias})
    problemas: list[str] = []
    pares: set[tuple[dt.date, dt.date]] = set()
    secuencia: list[tuple[dt.date, dt.date, dt.date]] = []  # (mes, ini, fin) de celdas propias
    for o in ocurrencias:
        if o.compartida:
            continue
        if o.tipo_fin == "INDET" or o.tipo_ini == "INDET":
            problemas.append("FECHA_INDETERMINADA")
        if o.tipo_ini == "INVALIDA" or o.tipo_fin == "INVALIDA":
            problemas.append("FECHA_NO_INTERPRETABLE")
        if (o.tipo_ini == "VACIA") != (o.tipo_fin == "VACIA"):
            problemas.append("FECHA_INCOMPLETA")
        # Ambas vacias: el bloque solo aporta presencia (no es un problema).
        if o.ini and o.fin:
            if o.ini > o.fin:
                problemas.append("INICIO_POSTERIOR_A_FIN")
            else:
                pares.add((o.ini, o.fin))
                secuencia.append((o.mes, o.ini, o.fin))
    if not pares:
        if any(o.compartida for o in ocurrencias):
            problemas.append("FECHAS_SOLO_EN_CELDA_COMPARTIDA")
        problemas.append("SIN_FECHAS_DETERMINABLES")
        return ResultadoFechas([], presencia, sorted(set(problemas)))

    clusters: list[list[tuple[dt.date, dt.date]]] = []
    for p in sorted(pares):
        if clusters and p[0] <= max(q[1] for q in clusters[-1]) + dt.timedelta(days=1):
            clusters[-1].append(p)
        else:
            clusters.append([p])
    periodos = []
    for cl in clusters:
        ini, fin = min(p[0] for p in cl), max(p[1] for p in cl)
        solapados = any(
            a != b and a[0] <= b[1] and b[0] <= a[1] for k, a in enumerate(cl) for b in cl[k + 1:]
        )
        if len(cl) == 1:
            regla = "RANGO_UNICO"
        elif not solapados:
            regla = "BLOQUES_CONTIGUOS_CONSOLIDADOS"
        elif es_extension_progresiva(cl, secuencia):
            regla = "EXTENSION_FIN_PROGRESIVA"
        else:
            regla = "RANGOS_SOLAPADOS"
        per = Periodo(ini=ini, fin=fin, regla=regla, pares=cl)
        if regla == "RANGOS_SOLAPADOS":
            per.dudas.append("RANGOS_DISTINTOS_SOLAPADOS")
        sin_presencia = [m for m in meses_entre(ini, fin) if m in ventana and m not in presencia]
        if sin_presencia:
            per.dudas.append("MESES_SIN_PRESENCIA:" + fmt_meses(sin_presencia))
        periodos.append(per)
    fuera = [m for m in presencia if not any(_intersecta_mes(p.ini, p.fin, m) for p in periodos)]
    if fuera:
        problemas.append("PRESENCIA_FUERA_DE_RANGO:" + fmt_meses(fuera))
    return ResultadoFechas(periodos, presencia, sorted(set(problemas)))


def es_extension_progresiva(
    cluster: list[tuple[dt.date, dt.date]], secuencia: list[tuple[dt.date, dt.date, dt.date]]
) -> bool:
    """Mismo INICIO en todos los pares y FIN no decreciente bloque a bloque.

    Patron tipico de la grilla: cada mes repite INICIO y actualiza FIN cuando
    la campana se extiende; el bloque mas reciente refleja el FIN vigente.
    Dentro de un mismo mes todas las celdas deben decir lo mismo.
    """
    if len({p[0] for p in cluster}) != 1:
        return False
    miembros = set(cluster)
    por_mes: dict[dt.date, set[dt.date]] = defaultdict(set)
    for mes, ini, fin in secuencia:
        if (ini, fin) in miembros:
            por_mes[mes].add(fin)
    if any(len(f) > 1 for f in por_mes.values()):
        return False
    fines = [next(iter(por_mes[m])) for m in sorted(por_mes)]
    return all(a <= b for a, b in zip(fines, fines[1:]))


# ---------------------------------------------------------------------------
# CAMPANAS: indice y deduplicacion
# ---------------------------------------------------------------------------

@dataclass
class FilaCampana:
    carga_id: str
    ini: dt.date | None
    fin: dt.date | None
    indefinida: bool


def _a_fecha(v: Any) -> dt.date | None:
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.date()
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def _a_ot(v: Any) -> int | None:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if pd.isna(f) or not f.is_integer():
        return None
    return int(f)


class IndiceCampanas:
    def __init__(self, campanas: pd.DataFrame):
        self.por_ot_elem: dict[tuple[int, str], list[FilaCampana]] = defaultdict(list)
        self.elems_por_ot: dict[int, set[str]] = defaultdict(set)
        self.atributos_por_ot: dict[int, set[tuple[str, ...]]] = defaultdict(set)
        for r in campanas.to_dict("records"):
            ot = _a_ot(r.get("IDCampaña"))
            if ot is None:
                continue
            eid = texto(r.get("ElementoID"))
            self.por_ot_elem[(ot, eid)].append(FilaCampana(
                carga_id=texto(r.get("CargaID")), ini=_a_fecha(r.get("FechaInicio")),
                fin=_a_fecha(r.get("FechaFin")), indefinida=texto(r.get("FechaIndefinida")) == "Si",
            ))
            self.elems_por_ot[ot].add(eid)
            self.atributos_por_ot[ot].add(tuple(texto(r.get(c)) for c in ("Campaña", "Cliente", "Marca", "Agencia", "Proveedor")))


def _intervalos(filas: list[FilaCampana]) -> list[tuple[dt.date, dt.date]]:
    out = []
    for f in filas:
        if f.ini is None:
            continue
        fin = f.fin or (dt.date.max if f.indefinida else None)
        if fin is None:
            continue
        out.append((f.ini, fin))
    out.sort()
    merged: list[list[dt.date]] = []
    for a, b in out:
        if merged and a <= merged[-1][1] + dt.timedelta(days=1):
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    return [(a, b) for a, b in merged]


def tramos_no_cubiertos(
    ini: dt.date, fin: dt.date, ivs: list[tuple[dt.date, dt.date]]
) -> list[tuple[dt.date, dt.date]]:
    """Partes de [ini, fin] que no cubre ninguno de los intervalos (fusionados)."""
    out, cursor = [], ini
    for a, b in ivs:
        if b < cursor or a > fin:
            continue
        if a > cursor:
            out.append((cursor, a - dt.timedelta(days=1)))
        cursor = max(cursor, b + dt.timedelta(days=1)) if b < dt.date.max else dt.date.max
        if cursor > fin:
            break
    if cursor <= fin:
        out.append((cursor, fin))
    return out


@dataclass
class ResultadoDedup:
    resultado: str  # NO_EXISTE, EXACTO, CUBIERTO, PARCIAL, OTRAS_FECHAS
    carga_ids: str = ""
    no_cubierto: list[tuple[dt.date, dt.date]] = field(default_factory=list)


def comparar_con_campanas(
    filas: list[FilaCampana], ini: dt.date | None, fin: dt.date | None, presencia: list[dt.date]
) -> ResultadoDedup:
    if not filas:
        return ResultadoDedup("NO_EXISTE")
    ids = ",".join(sorted(f.carga_id for f in filas))
    if ini and fin and any(f.ini == ini and f.fin == fin for f in filas):
        return ResultadoDedup("EXACTO", ids)
    ivs = _intervalos(filas)
    if ini and fin:
        faltan = tramos_no_cubiertos(ini, fin, ivs)
        if not faltan:
            return ResultadoDedup("CUBIERTO", ids)
        if faltan != [(ini, fin)]:
            return ResultadoDedup("PARCIAL", ids, faltan)
        return ResultadoDedup("OTRAS_FECHAS", ids, faltan)
    meses = [m for m in presencia if m.year == ANIO_ALCANCE] or presencia
    cubiertos = [m for m in meses if any(a <= fin_mes(m) and m <= b for a, b in ivs)]
    if meses and len(cubiertos) == len(meses):
        return ResultadoDedup("CUBIERTO", ids)
    if cubiertos:
        return ResultadoDedup("PARCIAL", ids)
    return ResultadoDedup("OTRAS_FECHAS", ids)


def clave_negocio(ot: int, eid: str, ini: dt.date, fin: dt.date) -> str:
    """Mismo formato que la migracion historica: OT|Elemento|Ini|Fin|HoraIni|HoraFin."""
    return f"{ot}|{eid}|{ini.isoformat()}|{fin.isoformat()}||"


# ---------------------------------------------------------------------------
# Consolidacion y clasificacion
# ---------------------------------------------------------------------------

def _uniq(vals: Iterable[str]) -> list[str]:
    out: list[str] = []
    for v in vals:
        if v and v not in out:
            out.append(v)
    return out


def consolidar(
    ocurrencias: list[Ocurrencia],
    crosswalk: pd.DataFrame,
    idx_m: IndiceMaestro,
    idx_c: IndiceCampanas,
    ventanas: dict[str, list[dt.date]],
) -> pd.DataFrame:
    xw = {(r["HojaOrigen"], r["ElementoOrigen"]): r for r in crosswalk.to_dict("records")}

    grupos: dict[tuple, list[Ocurrencia]] = defaultdict(list)
    for o in ocurrencias:
        if not o.codigo:
            grupos[("FILA_SIN_CODIGO", o.hoja, o.fila, o.token.numero if o.token else o.texto_no_ot)].append(o)
            continue
        x = xw.get((o.hoja, o.codigo), {})
        elem_key = x["ElementoID_OCU26"] if x.get("Confianza") == "ALTA" else f"{o.hoja}|{o.codigo}"
        if o.token is None:
            grupos[("SIN_OT", elem_key, o.texto_no_ot or o.ot_raw or "PAUTA:" + o.pauta.upper())].append(o)
        else:
            grupos[("OT", elem_key, o.token.numero, o.token.bonificada)].append(o)

    registros = []
    for key, occ in grupos.items():
        tipo = key[0]
        hojas = _uniq(o.hoja for o in occ)
        codigos = _uniq(o.codigo for o in occ)
        x = xw.get((occ[0].hoja, occ[0].codigo), {}) if occ[0].codigo else {}
        alcance = _uniq(o.alcance for o in occ)
        ventana = set().union(*(ventanas[h] for h in hojas))
        base = {
            "TRZ_TipoGrupo": tipo,
            "TRZ_HojaOrigen": " + ".join(hojas),
            "TRZ_ElementoOrigen": " + ".join(codigos),
            "TRZ_FilasOrigen": ", ".join(str(f) for f in sorted({o.fila for o in occ})),
            "TRZ_RefCeldas": ", ".join(sorted({o.ref for o in occ}, key=lambda s: (s.split("!")[0], int(re.sub(r"\D", "", s.split("!")[1]) or 0), s))),
            "TRZ_OT_Raw": " || ".join(_uniq(o.ot_raw for o in occ)),
            "TRZ_PrefijoOT_Origen": ", ".join(_uniq(o.token.prefijo for o in occ if o.token)),
            "TRZ_Pauta_Origen": " || ".join(_uniq(o.pauta for o in occ)),
            "TRZ_Obs_Origen": " || ".join(_uniq(o.obs for o in occ)),
            "TRZ_CeldaCompartida": "SI" if any(o.compartida for o in occ) else "NO",
            "TRZ_UbicacionOrigen": " || ".join(_uniq(o.attrs.get("Ubic", "") for o in occ)),
            "TRZ_DescripcionOrigen": " || ".join(_uniq(o.attrs.get("Descripcion", "") or o.attrs.get("Disp", "") for o in occ))[:500],
            "TRZ_CircuitoOrigen": circuito_origen(occ[0].hoja, occ[0].attrs) if occ[0].codigo else "",
            "TRZ_MedioOrigen": HOJAS_PROCESAR.get(occ[0].hoja) or "",
            "TRZ_AlcanceOrigen": " + ".join(alcance),
            "TRZ_Ocurrencias": len(occ),
            "TRZ_MetodoMatch": x.get("MetodoMatch", "") if x else "NO_MAPEADO_FUERA_DE_ALCANCE_O_SIN_CODIGO",
            "TRZ_ConfianzaMatch": x.get("Confianza", ""),
            "TRZ_CandidatosElemento": x.get("Candidatos", ""),
        }
        ot = key[2] if tipo == "OT" else (key[3] if tipo == "FILA_SIN_CODIGO" and isinstance(key[3], int) else None)
        bonificada = bool(key[3]) if tipo == "OT" else False
        eid = x.get("ElementoID_OCU26", "") if x.get("Confianza") == "ALTA" else ""
        fechas = reconstruir_periodos(occ, ventana)
        periodos: list[Periodo | None] = fechas.periodos or [None]
        regs_grupo = []
        for per in periodos:
            occ_p = occ if per is None else [o for o in occ if _intersecta_mes(per.ini, per.fin, o.mes)] or occ
            pres = sorted({o.mes for o in occ_p}) if per is not None else fechas.presencia
            reg = dict(base)
            reg.update({
                "IDCampaña": ot,
                "ElementoID": eid,
                "FechaInicio": per.ini if per else None,
                "FechaFin": per.fin if per else None,
                "TRZ_OT_Bonificada": "SI" if bonificada else "NO",
                "TRZ_MesesPresencia": fmt_meses(pres),
                "TRZ_ReglaFechas": per.regla if per else "NO_DETERMINABLE",
                "TRZ_ParesInicioFin": "; ".join(f"{a}→{b}" for a, b in (per.pares if per else [])),
                "TRZ_ProblemasFechas": " | ".join(fechas.problemas + (per.dudas if per else [])),
                "TRZ_PeriodosDisjuntos": "SI" if fechas.disjuntos else "NO",
            })
            estado, motivos, dd = clasificar(reg, tipo, ot, bonificada, eid, x, alcance, fechas, per, pres, idx_c)
            reg["Estado_Staging"] = estado
            reg["_motivos"] = motivos
            reg["TRZ_Dedup"] = dd.resultado if dd else ""
            reg["TRZ_CargaID_Existentes"] = dd.carga_ids if dd else ""
            reg["TRZ_TramosNoCubiertosEnCAMPANAS"] = "; ".join(f"{a}→{b}" for a, b in dd.no_cubierto) if dd else ""
            otros = sorted(idx_c.elems_por_ot.get(ot, set()) - {eid}) if ot is not None else []
            reg["TRZ_OT_en_CAMPANAS"] = "SI" if ot is not None and ot in idx_c.elems_por_ot else "NO"
            reg["TRZ_OT_otros_ElementoID"] = ", ".join(otros[:20]) + (" ..." if len(otros) > 20 else "")
            atrs = idx_c.atributos_por_ot.get(ot, set()) if ot is not None else set()
            reg["TRZ_Ref_Atributos_OT_CAMPANAS"] = (
                " / ".join(next(iter(atrs))) if len(atrs) == 1 else (f"{len(atrs)} variantes" if atrs else "")
            )
            regs_grupo.append(reg)
        promover_periodos_disjuntos(regs_grupo)
        for reg in regs_grupo:
            reg["Motivo"] = " | ".join(reg.pop("_motivos"))
        registros.extend(regs_grupo)
    return pd.DataFrame(registros)


MOTIVO_SOLO_OTRAS_FECHAS = "OT_ELEMENTO_EXISTE_CON_OTRAS_FECHAS"


def promover_periodos_disjuntos(regs_grupo: list[dict]) -> None:
    """Un periodo disjunto de una OT cuyo otro periodo ya esta en CAMPANAS es
    un nuevo tramo (hueco real), no una correccion de fechas -> IMPORTAR.

    Solo si ese era su UNICO motivo de revision."""
    if len(regs_grupo) < 2 or not any(r["Estado_Staging"] == "YA_EXISTE" for r in regs_grupo):
        return
    for r in regs_grupo:
        if r["Estado_Staging"] == "REVISAR" and r["_motivos"] == [MOTIVO_SOLO_OTRAS_FECHAS]:
            r["Estado_Staging"] = "IMPORTAR"
            r["_motivos"] = ["NUEVO_PERIODO_DISJUNTO_DE_OT_YA_CARGADA"]


def clasificar(reg, tipo, ot, bonificada, eid, x, alcance, fechas, per, pres, idx_c) -> tuple[str, list[str], ResultadoDedup | None]:
    motivos: list[str] = []
    # 1) Fuera de alcance -> NO_IMPORTAR
    excl = [a for a in alcance if a not in ("EN_ALCANCE", "FILA_SIN_CODIGO")]
    if excl:
        return "NO_IMPORTAR", [f"FUERA_DE_ALCANCE_{a}" for a in excl], None
    if tipo == "SIN_OT":
        txt = sin_acentos(f"{reg['TRZ_OT_Raw']} {reg['TRZ_Pauta_Origen']}").upper()
        if PATRON_NO_COMERCIAL.search(txt):
            return "NO_IMPORTAR", ["NO_COMERCIAL_SIN_OT"], None
    anios = {m.year for m in pres}
    en_2026 = per is not None and per.ini <= dt.date(ANIO_ALCANCE, 12, 31) and per.fin >= dt.date(ANIO_ALCANCE, 1, 1)
    if ANIO_ALCANCE not in anios and not en_2026:
        return "NO_IMPORTAR", ["FUERA_DE_2026"], None
    if tipo == "FILA_SIN_CODIGO":
        return "REVISAR", ["ESTRUCTURA_FILA_SIN_CODIGO_ELEMENTO"], None
    if tipo == "SIN_OT":
        return "REVISAR", ["SIN_OT_NUMERICA"], None

    # 2) Deduplicacion (solo con elemento resuelto)
    dd: ResultadoDedup | None = None
    if eid:
        dd = comparar_con_campanas(idx_c.por_ot_elem.get((ot, eid), []),
                                   per.ini if per else None, per.fin if per else None, pres)
        if dd.resultado in ("EXACTO", "CUBIERTO"):
            m = ["YA_EXISTE_EXACTO" if dd.resultado == "EXACTO" else "YA_EXISTE_CUBIERTO_FECHAS_DISTINTAS"]
            if per is None:
                m.append("FECHAS_ORIGEN_NO_DETERMINABLES")
            return "YA_EXISTE", m, dd
    else:
        cands = [c for c in texto(x.get("Candidatos", "")).split(", ") if c] if x else []
        hits = [c for c in cands if (ot, c) in idx_c.por_ot_elem]
        if hits:
            motivos.append("OT_EXISTE_EN_CAMPANAS_CON_CANDIDATO:" + ",".join(hits))

    # 3) Problemas -> REVISAR
    if not eid:
        conf = x.get("Confianza", "BAJA") if x else "BAJA"
        motivos.insert(0, f"ELEMENTO_NO_INEQUIVOCO_{conf}")
    if bonificada:
        motivos.append("OT_BONIFICADA_PREFIJO_B")
    if per is None:
        motivos.append("FECHAS_NO_DETERMINABLES")
    elif fechas.problemas or per.dudas:
        motivos.append("FECHAS_DUDOSAS")
    if dd and dd.resultado == "PARCIAL":
        motivos.append("EXISTE_EN_CAMPANAS_COBERTURA_PARCIAL")
    if dd and dd.resultado == "OTRAS_FECHAS":
        motivos.append(MOTIVO_SOLO_OTRAS_FECHAS)
    if per is not None and not en_2026:
        motivos.append("PERIODO_FUERA_DE_2026_CON_PRESENCIA_2026")
    if motivos:
        return "REVISAR", motivos, dd
    return "IMPORTAR", ["NUEVA_ASIGNACION"], dd


# ---------------------------------------------------------------------------
# Salidas
# ---------------------------------------------------------------------------

TRZ_ORDEN = [
    "StagingID", "Estado_Staging", "Motivo", "IDCampaña", "ElementoID", "FechaInicio", "FechaFin",
    "TRZ_Medio_OCU26", "TRZ_CircuitoDashboard_OCU26", "TRZ_HojaOrigen", "TRZ_ElementoOrigen", "TRZ_MedioOrigen", "TRZ_CircuitoOrigen", "TRZ_UbicacionOrigen",
    "TRZ_DescripcionOrigen", "TRZ_MetodoMatch", "TRZ_ConfianzaMatch", "TRZ_CandidatosElemento",
    "TRZ_OT_Raw", "TRZ_PrefijoOT_Origen", "TRZ_OT_Bonificada", "TRZ_Pauta_Origen", "TRZ_Obs_Origen",
    "TRZ_MesesPresencia", "TRZ_ReglaFechas", "TRZ_ParesInicioFin", "TRZ_ProblemasFechas",
    "TRZ_PeriodosDisjuntos", "TRZ_CeldaCompartida", "TRZ_Dedup", "TRZ_CargaID_Existentes",
    "TRZ_TramosNoCubiertosEnCAMPANAS",
    "TRZ_OT_en_CAMPANAS", "TRZ_OT_otros_ElementoID", "TRZ_Ref_Atributos_OT_CAMPANAS",
    "TRZ_AlcanceOrigen", "TRZ_TipoGrupo", "TRZ_Ocurrencias", "TRZ_FilasOrigen", "TRZ_RefCeldas",
]


def preparar_importar(df: pd.DataFrame, medio_por_id: dict[str, str]) -> pd.DataFrame:
    """Columnas CAMPANAS exactas (orden vigente) + columnas TRZ_ separadas."""
    imp = df[df["Estado_Staging"] == "IMPORTAR"].copy()
    destino = pd.DataFrame(index=imp.index, columns=CAMPANAS_HEADERS, dtype=object)
    destino["IDCampaña"] = imp["IDCampaña"].astype(int)
    destino["ElementoID"] = imp["ElementoID"]
    destino["FechaInicio"] = imp["FechaInicio"]
    destino["FechaFin"] = imp["FechaFin"]
    destino["ClaveNegocio"] = [clave_negocio(int(o), e, a, b) for o, e, a, b in
                               zip(imp["IDCampaña"], imp["ElementoID"], imp["FechaInicio"], imp["FechaFin"])]
    destino["TipoCargaDeclarado"] = imp["ElementoID"].map(medio_por_id)
    destino["FechaIndefinida"] = "No"
    # Campaña = PAUTA literal de origen solo si es unica para la asignacion.
    destino["Campaña"] = [p if p and "||" not in p else None for p in imp["TRZ_Pauta_Origen"]]
    trz = imp[[c for c in TRZ_ORDEN if c in imp.columns and c not in ("IDCampaña", "ElementoID", "FechaInicio", "FechaFin")]]
    out = pd.concat([destino, trz], axis=1)
    if out["ClaveNegocio"].duplicated().any():
        dups = out.loc[out["ClaveNegocio"].duplicated(keep=False), "ClaveNegocio"].tolist()
        raise StagingError(f"ClaveNegocio duplicada dentro de IMPORTAR: {dups[:5]}")
    return out


def construir_staging(fuente: Path, base: Path) -> dict[str, Any]:
    ocurrencias, ventanas, no_procesadas = leer_fuente(fuente)
    maestro = pd.read_excel(base, sheet_name="MAESTRO_ELEMENTOS")
    campanas = pd.read_excel(base, sheet_name="CAMPANAS")
    if list(campanas.columns) != CAMPANAS_HEADERS:
        raise StagingError("CAMPANAS: encabezados distintos del esquema vigente")
    idx_m = IndiceMaestro(maestro)
    idx_c = IndiceCampanas(campanas)
    xw = construir_crosswalk(ocurrencias, idx_m)
    asign = consolidar(ocurrencias, xw, idx_m, idx_c, ventanas)
    orden = {e: k for k, e in enumerate(ESTADOS)}
    asign["_o"] = asign["Estado_Staging"].map(orden)
    asign["_ini"] = asign["FechaInicio"].map(lambda d: d or dt.date.min)
    asign["_ot"] = asign["IDCampaña"].map(lambda v: -1 if v is None or pd.isna(v) else int(v))
    asign = asign.sort_values(["_o", "TRZ_HojaOrigen", "_ot", "ElementoID", "TRZ_ElementoOrigen", "_ini", "TRZ_RefCeldas"],
                              kind="stable").drop(columns=["_o", "_ini", "_ot"]).reset_index(drop=True)
    asign.insert(0, "StagingID", [f"STG2B-{k + 1:06d}" for k in range(len(asign))])
    asign["IDCampaña"] = asign["IDCampaña"].map(lambda v: None if v is None or pd.isna(v) else int(v)).astype(object)
    medio_por_id = {e: texto(r.get("Medio")) for e, r in idx_m.por_id.items()}
    circ_por_id = {e: texto(r.get("CircuitoDashboard")) for e, r in idx_m.por_id.items()}
    asign["TRZ_Medio_OCU26"] = asign["ElementoID"].map(lambda e: medio_por_id.get(e, "") if e else "")
    asign["TRZ_CircuitoDashboard_OCU26"] = asign["ElementoID"].map(lambda e: circ_por_id.get(e, "") if e else "")
    hojas = {
        "IMPORTAR": preparar_importar(asign, medio_por_id),
        "YA_EXISTE": asign[asign["Estado_Staging"] == "YA_EXISTE"][TRZ_ORDEN],
        "REVISAR": asign[asign["Estado_Staging"] == "REVISAR"][TRZ_ORDEN],
        "NO_IMPORTAR": asign[asign["Estado_Staging"] == "NO_IMPORTAR"][TRZ_ORDEN],
    }
    return {
        "ocurrencias": ocurrencias, "ventanas": ventanas, "no_procesadas": no_procesadas,
        "crosswalk": xw, "asignaciones": asign, "hojas": hojas, "campanas": campanas,
    }


# ---------------------------------------------------------------------------
# Resumen
# ---------------------------------------------------------------------------

def meses_activos(row: pd.Series) -> list[dt.date]:
    if row["FechaInicio"] and row["FechaFin"]:
        return meses_entre(row["FechaInicio"], row["FechaFin"])
    return [dt.date(int(m[:4]), int(m[5:7]), 1) for m in texto(row["TRZ_MesesPresencia"]).split(", ") if m]


def _ots(df: pd.DataFrame) -> set[int]:
    return {int(v) for v in df["IDCampaña"] if v is not None and not pd.isna(v)}


def _conteo_estados(df: pd.DataFrame) -> dict[str, int]:
    vc = df["Estado_Staging"].value_counts()
    return {e: int(vc.get(e, 0)) for e in ESTADOS}


def calcular_resumen(res: dict[str, Any]) -> dict[str, Any]:
    a: pd.DataFrame = res["asignaciones"]
    xw: pd.DataFrame = res["crosswalk"]
    occ: list[Ocurrencia] = res["ocurrencias"]
    campanas_ots = {o for o in (_a_ot(v) for v in res["campanas"]["IDCampaña"]) if o is not None}

    def fila(nombre: str, sub_a: pd.DataFrame, sub_occ: list[Ocurrencia], sub_xw: pd.DataFrame) -> dict[str, Any]:
        ce = _conteo_estados(sub_a)
        imp_ots = _ots(sub_a[sub_a["Estado_Staging"] == "IMPORTAR"])
        return {
            "Ambito": nombre,
            "Ocurrencias celda x OT": len(sub_occ),
            "Celdas-bloque con datos": len({(o.hoja, o.ref) for o in sub_occ}),
            "Asignaciones consolidadas": len(sub_a),
            "OT distintas": len(_ots(sub_a)),
            **ce,
            "OT distintas IMPORTAR": len(imp_ots),
            "OT nuevas (inexistentes en CAMPANAS)": len(imp_ots - campanas_ots),
            "Elementos historicos distintos (en alcance)": len(sub_xw),
            "Elementos mapeados ALTA": int((sub_xw["Confianza"] == "ALTA").sum()),
            "ElementoID OCU26 distintos (ALTA)": sub_xw.loc[sub_xw["Confianza"] == "ALTA", "ElementoID_OCU26"].nunique(),
            "Elementos sin match (BAJA)": int((sub_xw["Confianza"] == "BAJA").sum()),
            "Matches ambiguos (MEDIA)": int((sub_xw["Confianza"] == "MEDIA").sum()),
            "Elementos fuera de alcance (no mapeados)": len({(o.hoja, o.codigo) for o in sub_occ
                                                           if o.codigo and o.alcance != "EN_ALCANCE"}),
        }

    general = [fila("TOTAL", a, occ, xw)]
    for hoja in HOJAS_PROCESAR:
        general.append(fila(hoja, a[a["TRZ_HojaOrigen"].str.contains(re.escape(hoja))],
                            [o for o in occ if o.hoja == hoja], xw[xw["HojaOrigen"] == hoja]))
    general_df = pd.DataFrame(general)

    # Por mes (2026): OT distintas y asignaciones activas por estado.
    filas_mes = []
    exploded = [(m, r["Estado_Staging"], r["IDCampaña"]) for _, r in a.iterrows() for m in meses_activos(r)]
    ex = pd.DataFrame(exploded, columns=["Mes", "Estado", "OT"])
    ex = ex[ex["Mes"].map(lambda m: m.year == ANIO_ALCANCE)]
    for mes in [dt.date(ANIO_ALCANCE, k, 1) for k in range(1, 13)]:
        sub = ex[ex["Mes"] == mes]
        d = {"Mes": mes.strftime("%Y-%m")}
        for e in ESTADOS:
            se = sub[sub["Estado"] == e]
            d[f"OT {e}"] = se["OT"].dropna().nunique()
            d[f"Asig {e}"] = len(se)
        filas_mes.append(d)
    por_mes = pd.DataFrame(filas_mes).rename(columns={"OT IMPORTAR": "OT nuevas (IMPORTAR)", "OT YA_EXISTE": "OT ya existentes"})

    def por_dim(col: str, etiqueta: str) -> pd.DataFrame:
        b = a.copy()
        b[etiqueta] = b[col].where(b[col] != "", "SIN_MAPEO")
        t = b.pivot_table(index=etiqueta, columns="Estado_Staging", values="StagingID", aggfunc="count", fill_value=0)
        t = t.reindex(columns=list(ESTADOS), fill_value=0)
        t["OT distintas"] = b.groupby(etiqueta)["IDCampaña"].apply(lambda s: s.dropna().nunique())
        t["OT nuevas (IMPORTAR)"] = b[b["Estado_Staging"] == "IMPORTAR"].groupby(etiqueta)["IDCampaña"].nunique()
        return t.fillna(0).astype(int).reset_index()

    b = a.assign(_m=[m or o for m, o in zip(a["TRZ_Medio_OCU26"], a["TRZ_MedioOrigen"])])
    b["Medio"] = b["_m"].where(b["_m"] != "", "MIXTO/SIN_MAPEO")
    por_medio = b.pivot_table(index="Medio", columns="Estado_Staging", values="StagingID", aggfunc="count",
                              fill_value=0).reindex(columns=list(ESTADOS), fill_value=0)
    por_medio["OT distintas"] = b.groupby("Medio")["IDCampaña"].apply(lambda s: s.dropna().nunique())
    por_medio = por_medio.fillna(0).astype(int).reset_index()
    por_circuito = por_dim("TRZ_CircuitoDashboard_OCU26", "CircuitoDashboard_OCU26")

    rv = a[a["Estado_Staging"] == "REVISAR"]
    motivos = Counter(re.sub(r":.*", "", m) for ms in rv["Motivo"] for m in ms.split(" | ") if m)
    motivos_df = pd.DataFrame(motivos.most_common(), columns=["Motivo REVISAR", "Asignaciones"])
    probs = Counter(re.sub(r":.*", "", p) for ps in rv["TRZ_ProblemasFechas"] for p in ps.split(" | ") if p)
    probs_df = pd.DataFrame(probs.most_common(), columns=["Problema de fechas (REVISAR)", "Asignaciones"])
    no_imp = a[a["Estado_Staging"] == "NO_IMPORTAR"]
    motivos_no = Counter(m for ms in no_imp["Motivo"] for m in ms.split(" | ") if m)
    motivos_no_df = pd.DataFrame(motivos_no.most_common(), columns=["Motivo NO_IMPORTAR", "Asignaciones"])

    dudosas = rv[rv["Motivo"].str.contains("FECHAS_DUDOSAS|FECHAS_NO_DETERMINABLES")]
    ots_dudosas = sorted(_ots(dudosas))
    disj = a[(a["TRZ_PeriodosDisjuntos"] == "SI") & a["IDCampaña"].notna()]
    ots_disjuntas = sorted(_ots(disj))
    no_encontrados = xw[xw["Confianza"] == "BAJA"][["HojaOrigen", "ElementoOrigen", "UbicacionOrigen", "DescripcionOrigen", "Observacion"]]
    ambiguos = xw[xw["Confianza"] == "MEDIA"][["HojaOrigen", "ElementoOrigen", "MetodoMatch", "Candidatos", "Observacion"]]
    reglas = (xw.groupby(["MetodoMatch", "Confianza"])
              .agg(Elementos=("ElementoOrigen", "size"),
                   Ejemplo=("ElementoOrigen", lambda s: s.iloc[0]),
                   Ejemplo_OCU26=("ElementoID_OCU26", lambda s: s.iloc[0]))
              .reset_index())

    hojas_proc = set(HOJAS_PROCESAR)
    trip_rows = [o for o in occ if o.hoja == "TRIPSTORE Y LS D"]
    imp_alc = set(" + ".join(a.loc[a["Estado_Staging"].isin(["IMPORTAR", "YA_EXISTE", "REVISAR"]), "TRZ_AlcanceOrigen"]).split(" + "))
    confirmaciones = {
        "TRIPSTORE procesado": "SI" if any(o.alcance == "EN_ALCANCE" for o in trip_rows) else "NO",
        "LS D excluido": "SI" if "LONDON/LS" not in imp_alc else "NO",
        "LS F excluido": "SI" if "LS F" not in hojas_proc else "NO",
        "LONDON excluido": "SI" if not any("LONDON" in x for x in imp_alc) else "NO",
        "YPF excluido": "SI" if "YPF" not in hojas_proc and "YPF" not in imp_alc else "NO",
        "CENCOMEDIA excluido": "SI" if "CENCOMEDIA" not in hojas_proc and "CENCOMEDIA" not in imp_alc else "NO",
    }
    return {
        "general": general_df, "por_mes": por_mes, "por_medio": por_medio, "por_circuito": por_circuito,
        "motivos_revisar": motivos_df, "problemas_fechas": probs_df, "motivos_no_importar": motivos_no_df,
        "ots_dudosas": ots_dudosas, "ots_disjuntas": ots_disjuntas, "no_encontrados": no_encontrados,
        "ambiguos": ambiguos, "reglas": reglas, "confirmaciones": confirmaciones,
        "no_procesadas": res["no_procesadas"],
    }


# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------

def _escribir_df(ws, df: pd.DataFrame, fila: int, destino: set[str] | None = None) -> int:
    from openpyxl.styles import Font, PatternFill

    bold = Font(bold=True)
    fill_dest = PatternFill("solid", fgColor="DDEBF7")
    fill_trz = PatternFill("solid", fgColor="EDEDED")
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=fila, column=j, value=str(col))
        c.font = bold
        if destino is not None:
            c.fill = fill_dest if col in destino else fill_trz
    for i, row in enumerate(df.itertuples(index=False), fila + 1):
        for j, v in enumerate(row, 1):
            if v is None or (isinstance(v, float) and pd.isna(v)):
                continue
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, (dt.date, dt.datetime)):
                c.number_format = "yyyy-mm-dd"
    return fila + len(df) + 1


def escribir_excel(path: Path, res: dict[str, Any], resumen: dict[str, Any], meta: list[tuple[str, str]]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre in ESTADOS:
        df = res["hojas"][nombre]
        ws = wb.create_sheet(nombre)
        _escribir_df(ws, df, 1, destino=set(CAMPANAS_HEADERS) if nombre == "IMPORTAR" else {"IDCampaña", "ElementoID", "FechaInicio", "FechaFin"})
        ws.freeze_panes = "B2"
        if len(df.columns):
            ws.auto_filter.ref = f"A1:{col_letra(len(df.columns) - 1)}{len(df) + 1}"
        for j, col in enumerate(df.columns):
            ws.column_dimensions[col_letra(j)].width = min(45, max(10, len(str(col)) + 2))
    xw = res["crosswalk"]
    ws = wb.create_sheet("CROSSWALK_ELEMENTOS")
    _escribir_df(ws, xw, 1)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{col_letra(len(xw.columns) - 1)}{len(xw) + 1}"
    for j, col in enumerate(xw.columns):
        ws.column_dimensions[col_letra(j)].width = min(45, max(12, len(str(col)) + 2))

    ws = wb.create_sheet("RESUMEN")
    from openpyxl.styles import Font
    titulo = Font(bold=True, size=12)
    f = 1
    bloques: list[tuple[str, pd.DataFrame]] = [
        ("METADATOS", pd.DataFrame(meta, columns=["Clave", "Valor"])),
        ("CONFIRMACIONES DE ALCANCE", pd.DataFrame(list(resumen["confirmaciones"].items()), columns=["Control", "Resultado"])),
        ("HOJAS NO PROCESADAS", pd.DataFrame(sorted(resumen["no_procesadas"].items()), columns=["Hoja", "Tratamiento"])),
        ("RESULTADOS GENERALES Y POR HOJA", resumen["general"]),
        ("CAMPANAS POR MES 2026 (OT distintas / asignaciones activas)", resumen["por_mes"]),
        ("CAMPANAS POR MEDIO", resumen["por_medio"]),
        ("CAMPANAS POR CIRCUITO (CircuitoDashboard OCU26)", resumen["por_circuito"]),
        ("MOTIVOS DE REVISAR", resumen["motivos_revisar"]),
        ("PROBLEMAS DE FECHAS EN REVISAR", resumen["problemas_fechas"]),
        ("MOTIVOS DE NO_IMPORTAR", resumen["motivos_no_importar"]),
        ("REGLAS DE NOMENCLATURA / METODOS DE MATCH", resumen["reglas"]),
        ("OTs CON FECHAS DUDOSAS O NO DETERMINABLES", pd.DataFrame({"OT": resumen["ots_dudosas"]})),
        ("OTs CON PERIODOS DISJUNTOS", pd.DataFrame({"OT": resumen["ots_disjuntas"]})),
        ("ELEMENTOS HISTORICOS NO ENCONTRADOS (BAJA)", resumen["no_encontrados"]),
        ("MATCHES AMBIGUOS (MEDIA)", resumen["ambiguos"]),
    ]
    for t, df in bloques:
        ws.cell(row=f, column=1, value=t).font = titulo
        f = _escribir_df(ws, df, f + 1) + 1
    ws.column_dimensions["A"].width = 48
    for j in range(1, 20):
        ws.column_dimensions[col_letra(j)].width = 18
    wb.save(path)


def escribir_csv(path: Path, importar: pd.DataFrame) -> None:
    df = importar[CAMPANAS_HEADERS].copy()
    for c in ("FechaInicio", "FechaFin"):
        df[c] = df[c].map(fmt_fecha)
    df.to_csv(path, index=False, encoding="utf-8-sig")


def _md_tabla(df: pd.DataFrame) -> str:
    if df.empty:
        return "_(sin filas)_\n"
    cols = [str(c) for c in df.columns]
    out = ["| " + " | ".join(cols) + " |", "|" + "|".join("---" for _ in cols) + "|"]
    for row in df.itertuples(index=False):
        out.append("| " + " | ".join(texto(v).replace("|", "/") for v in row) + " |")
    return "\n".join(out) + "\n"


def escribir_md(path: Path, res: dict[str, Any], resumen: dict[str, Any], meta: list[tuple[str, str]]) -> None:
    g = resumen["general"]
    tot = g.iloc[0]
    xw = res["crosswalk"]
    l: list[str] = []
    l.append("# Resumen de importacion historica OCU26 - Etapa 2B (staging, sin importar)\n")
    l.append(_md_tabla(pd.DataFrame(meta, columns=["Clave", "Valor"])))
    l.append("\n## Confirmaciones de alcance\n")
    l.append(_md_tabla(pd.DataFrame(list(resumen["confirmaciones"].items()), columns=["Control", "Resultado"])))
    l.append("\nHojas no procesadas:\n")
    l.append(_md_tabla(pd.DataFrame(sorted(resumen["no_procesadas"].items()), columns=["Hoja", "Tratamiento"])))
    l.append("\n## Resultado general\n")
    l.append(f"- Ocurrencias analizadas (celda x OT): **{tot['Ocurrencias celda x OT']}** "
             f"({tot['Celdas-bloque con datos']} celdas-bloque con datos)\n")
    l.append(f"- Asignaciones consolidadas (IDCampaña + elemento + periodo): **{tot['Asignaciones consolidadas']}**\n")
    l.append(f"- OT distintas: **{tot['OT distintas']}**\n")
    for e in ESTADOS:
        l.append(f"- {e}: **{tot[e]}**\n")
    l.append(f"- OT nuevas (IMPORTAR, inexistentes en CAMPANAS): **{tot['OT nuevas (inexistentes en CAMPANAS)']}**\n")
    l.append(f"- Elementos historicos distintos en alcance: **{tot['Elementos historicos distintos (en alcance)']}** "
             f"(ALTA {tot['Elementos mapeados ALTA']}, MEDIA {tot['Matches ambiguos (MEDIA)']}, BAJA {tot['Elementos sin match (BAJA)']})\n")
    l.append("\n## Por hoja\n")
    l.append(_md_tabla(g))
    l.append("\n## Campanas por mes 2026\n")
    l.append(_md_tabla(resumen["por_mes"]))
    l.append("\n## Campanas por medio\n")
    l.append(_md_tabla(resumen["por_medio"]))
    l.append("\n## Campanas por circuito\n")
    l.append(_md_tabla(resumen["por_circuito"]))
    l.append("\n## Motivos de REVISAR\n")
    l.append(_md_tabla(resumen["motivos_revisar"]))
    l.append("\n### Problemas de fechas en REVISAR\n")
    l.append(_md_tabla(resumen["problemas_fechas"]))
    l.append("\n## Motivos de NO_IMPORTAR\n")
    l.append(_md_tabla(resumen["motivos_no_importar"]))
    l.append("\n## Reglas de nomenclatura detectadas (crosswalk)\n")
    l.append(_md_tabla(resumen["reglas"]))
    un_uni = xw[xw["MetodoMatch"].str.contains("UN_UNI")]
    l.append(f"\n- UN->UNI confirmado contra maestro: {len(un_uni)} codigos historicos "
             f"(ej. {', '.join(f'{a} -> {b}' for a, b in un_uni[['ElementoOrigen', 'ElementoID_OCU26']].head(4).itertuples(index=False))}). "
             "No se aplico replace global: cada caso se confirmo contra MAESTRO_ELEMENTOS.\n")
    pled = xw[xw["MetodoMatch"] == "PLED_DESCRIPCION_MAESTRO"]
    eq = (pled.assign(raiz=pled["ElementoOrigen"].map(lambda c: _sin_sufijo_slot(c) or normalizar_codigo(c)))
          .groupby(["raiz", "ElementoID_OCU26", "UbicacionOrigen", "Ubicacion_OCU26"]).size().reset_index(name="Slots"))
    l.append("\n### Equivalencias PLED (codigo historico = Descripcion del maestro)\n")
    l.append(_md_tabla(eq))
    dp = xw[xw["MetodoMatch"] == "DESCRIPCION_Y_POSICION"]
    eqd = (dp.assign(raiz=dp["ElementoOrigen"].map(lambda c: _sin_sufijo_slot(c) or normalizar_codigo(c)))
           .groupby(["raiz", "ElementoID_OCU26", "DescripcionOrigen"]).size().reset_index(name="Slots"))
    l.append("\n### Equivalencias por Descripcion + posicion (UN/UNI-DOBLES -> UNI-TOTEMD)\n")
    l.append(_md_tabla(eqd))
    l.append("\n## Elementos historicos no encontrados (BAJA)\n")
    l.append(_md_tabla(resumen["no_encontrados"]))
    l.append("\n## Matches ambiguos (MEDIA)\n")
    l.append(_md_tabla(resumen["ambiguos"]))
    l.append(f"\n## OTs con fechas dudosas o no determinables ({len(resumen['ots_dudosas'])})\n\n")
    l.append(", ".join(str(o) for o in resumen["ots_dudosas"]) + "\n")
    l.append(f"\n## OTs con periodos disjuntos ({len(resumen['ots_disjuntas'])})\n\n")
    l.append(", ".join(str(o) for o in resumen["ots_disjuntas"]) + "\n")
    l.append(REGLAS_MD)
    path.write_text("".join(x if x.endswith("\n") else x + "\n" for x in l), encoding="utf-8")


REGLAS_MD = """
## Reglas aplicadas

- **Grano**: una asignacion = IDCampaña (OT) + ElementoID + periodo. Las filas
  de caras (estatico) o slots -Vn (digital) del mismo elemento se consolidan.
- **OT**: numero de la celda del bloque mensual. Prefijos (PUBLI, PORCO,
  SEGNO, ...) se conservan en `TRZ_PrefijoOT_Origen` (no se copian a
  Proveedor). `B <n>` = OT bonificada: REVISAR (no existe como IDCampaña en
  CAMPANAS).
- **Fechas**: pares INICIO/FIN identicos o contiguos se consolidan; pares sin
  contacto = periodos separados (huecos reales); mismo INICIO con FIN que
  crece bloque a bloque = extension (vale el ultimo FIN). Cualquier otro
  solape, mes del rango sin presencia de la OT, presencia fuera de rango,
  fecha vacia/indeterminada/no interpretable o fechas solo en celdas con
  varias OT -> REVISAR.
- **Dedup contra CAMPANAS** (IDCampaña + ElementoID): mismas fechas o rango
  contenido en filas existentes -> YA_EXISTE; cobertura parcial (p.ej.
  extension no cargada) u otras fechas -> REVISAR (no se modifica CAMPANAS),
  salvo periodo disjunto de una OT cuyo otro periodo ya esta cargado ->
  IMPORTAR.
- **IMPORTAR**: elemento ALTA + OT numerica + fechas inequivocas + periodo en
  2026 + ausente en CAMPANAS. Columnas destino = esquema CAMPANAS vigente.
  Derivadas con certeza: IDCampaña, ElementoID, FechaInicio, FechaFin,
  ClaveNegocio (misma formula que la migracion historica, verificada contra
  CAMPANAS), TipoCargaDeclarado (= Medio del maestro), FechaIndefinida = No,
  Campaña = PAUTA literal de origen (solo si es unica). Vacias a proposito:
  CargaID, FechaHoraCarga, UsuarioCarga, FuenteCarga, EstadoValidacion (se
  asignan al importar) y Cliente, Marca, Agencia, Proveedor, Estado, etc.
"""


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--fuente", required=True, help="OCUPACION_2026.xlsx")
    ap.add_argument("--base", default=None, help="OCU26_BASE_DATOS.xlsx (default: resolucion del pipeline)")
    ap.add_argument("--salida", required=True, help="Directorio de salida del staging")
    ap.add_argument("--sha-base", default=EXPECTED_BASE_SHA256, help="SHA-256 esperado de la base")
    args = ap.parse_args(argv)

    fuente = Path(args.fuente)
    base = resolve_input_path(args.base)
    for p in (fuente, base):
        if not p.is_file():
            print(f"ERROR: no existe {p}", file=sys.stderr)
            return 2
    sha_fuente, sha_base = sha256_file(fuente), sha256_file(base)
    print(f"FUENTE {fuente} sha256={sha_fuente}")
    print(f"BASE   {base} sha256={sha_base}")
    if sha_base != args.sha_base:
        print(f"ERROR: SHA de la base {sha_base} != esperado {args.sha_base}. Tarea detenida.", file=sys.stderr)
        return 2

    res = construir_staging(fuente, base)
    resumen = calcular_resumen(res)
    salida = Path(args.salida)
    salida.mkdir(parents=True, exist_ok=True)
    xlsx = salida / "OCU26_HISTORICO_STAGING.xlsx"
    csv = salida / "OCU26_HISTORICO_IMPORTAR.csv"
    md = salida / "RESUMEN_IMPORTACION_HISTORICA.md"

    sha_fuente_fin, sha_base_fin = sha256_file(fuente), sha256_file(base)
    meta = [
        ("Fuente", str(fuente)), ("SHA-256 fuente (inicio)", sha_fuente), ("SHA-256 fuente (fin)", sha_fuente_fin),
        ("Base OCU26", str(base)), ("SHA-256 base (inicio)", sha_base), ("SHA-256 base (fin)", sha_base_fin),
        ("SHA-256 base esperado", args.sha_base), ("Anio de alcance", str(ANIO_ALCANCE)),
        ("Generado", dt.datetime.now().strftime("%Y-%m-%d %H:%M")), ("Importacion realizada", "NO (solo staging)"),
    ]
    escribir_excel(xlsx, res, resumen, meta)
    escribir_csv(csv, res["hojas"]["IMPORTAR"])
    escribir_md(md, res, resumen, meta)
    if (sha256_file(fuente), sha256_file(base)) != (sha_fuente, sha_base):
        print("ERROR: un Excel fuente cambio durante el proceso", file=sys.stderr)
        return 3
    print(resumen["general"].to_string(index=False))
    print(f"Escritos: {xlsx}\n          {csv}\n          {md}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
